import mechanicalsoup
import urllib3
import pandas as pd
import csv
import re
from openpyxl import load_workbook
import urllib.request, json
isbns = []
invalid_link = []
# -- change --
wb = load_workbook(filename='src.xlsx')
ws = wb['Sheet1']
start = 1
end = 300
column = 'O'
# -- change --

for i in range(start, end + 1): # change 
    isbns.append(str(ws[column + str(i)].value).zfill(10)) # change
print('start at line ' + str(start) + '\nend at line ' + str(end))
print('all isbns: ' +str(isbns))
browser = mechanicalsoup.StatefulBrowser()

lst_dic = []
# firstTime = True
df = pd.DataFrame()
index = 0

def runDic(browser):
    diction = {}
    try:        
        diction['title'] = browser.get_current_page().select_one(
            '#bibdata').select_one('.title').text.replace('\n', '')
    except Exception:
        diction['title'] = None
    try:    
        diction["author"] = browser.get_current_page().select_one(
            '#bib-author-cell').text.replace('\n', '')
    except Exception:
        diction["author"] = None
    try:
        diction["publisher"] = browser.get_current_page().select_one(
            '#bib-publisher-cell').text.replace('\n', '')
    except Exception:
        diction["publisher"] = None
    try:    
        diction["edition_format"] = browser.get_current_page().select_one(
            '#bib-itemType-cell').text.replace('\n', '').replace('\xa0', '')
    except Exception:
        diction["edition_format"] = None
    try:    
        diction["summary"] = browser.get_current_page().select_one(
            '#bib-summary-cell').text.replace('\n', '').strip()
    except Exception:
        diction["summary"] = None
    try:    
        diction["subjects"] = browser.get_current_page().select_one(
            '#subject-terms').text.replace('\n', '').replace(' -- ', ' ')
    except Exception:
        diction["subjects"] = None
    try:    
        diction["genre"] = browser.get_current_page().select_one(
            '#details-genre').text.replace('\n', '')
    except Exception:
        diction["genre"] = None

    try:    
        diction["doctype"] = browser.get_current_page().select_one(
            '#details-doctype').text.replace('\n', '')
    except Exception:
        diction["doctype"] = None
    try:    
        diction["notes"] = browser.get_current_page().select_one(
            '#details-notes').text.replace('\n', '')
    except Exception:
        diction["notes"] = None
    try:    
        diction["ISBN"] = [browser.get_current_page().select_one(
            '#details-standardno').text.replace('\n', '').strip('ISBN:').replace(' ', ' ')]
    except Exception:
        diction["ISBN"] = None
    try:    
        diction["responsibility"] = browser.get_current_page(
        ).select_one('#details-respon').text.replace('\n', '')
    except Exception:
        diction["responsibility"] = None
    try:
        for i in browser.get_current_page().select_one('#details').select_one('div').select_one('table').select('tr'):
            if 'Material Type:' in i.select_one('th').text:
                diction["metType"] = i.select_one('td').text
    except Exception:
        diction["metType"] = None

    return diction


for link in isbns:
    index += 1
    
    with urllib.request.urlopen(f'http://xisbn.worldcat.org/webservices/xid/isbn/{link}?method=getEditions&fl=*&format=json') as url:
        data = json.loads(url.read().decode())
        try:
            # can use this directely
            lst = data[list(data.keys())[1]]
        except:
            invalid_link.append(link);
            continue
        dirUrl = lst[0]['url'][0]

        try:
            with browser.open(dirUrl):
                diction = runDic(browser)
                print('valid index ' + str(index))
                if (diction):
                    lst_dic.append(diction)
        except Exception as e:
            continue

df_valid = pd.DataFrame(lst_dic)
df_valid.to_csv('valid.csv', mode = 'a')
print('invalid isbns: ' + str(invalid_link))
df_invalid = pd.DataFrame({
    "invalid link": invalid_link
})

df_invalid.to_csv('invalid.csv', mode = 'a')
